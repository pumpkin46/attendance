"""Styled report exports (CSV / Excel / PDF) for the attendance range report,
plus the shared scaffolding (title band, KPI cards, table styling) reused by
the audit-trail and security-alert exports.

The export endpoint reuses the exact payload the page renders
(``report_service.attendance_range_report``) so a downloaded file always
matches what the user saw on screen. This module is pure presentation: it
turns those payloads into bytes.

Format conventions:
- CSV is machine-readable: raw minutes, ISO-8601 timestamps (localized to the
  viewer's offset), method columns, UTF-8 BOM so Excel detects the encoding.
- Excel and PDF are presentation documents: title band, KPI cards, a totals
  row, friendly durations ("7h 30m") and status colors. Excel sheets are
  print-ready (fit to width, repeated header row, page footer).
- Times are rendered in the viewer's timezone via ``tz_offset`` (the value of
  JavaScript's ``Date.getTimezoneOffset()``: minutes, UTC minus local).

openpyxl / reportlab are imported lazily inside the renderers (style objects
come from cached factories), so importing this module — and therefore app
startup — does not pay their import cost.
"""

from __future__ import annotations

import codecs
import csv
import io
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from functools import lru_cache
from types import SimpleNamespace
from typing import TYPE_CHECKING

from app.schemas.report import AttendanceRangeReport

if TYPE_CHECKING:  # heavy deps are runtime-lazy; keep them type-only here
    from openpyxl import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

# Shared palette (matches the app's slate/blue theme).
_NAVY = "1E293B"
_BLUE = "2563EB"
_SLATE = "64748B"
_CARD = "F8FAFC"
_ZEBRA = "F1F5F9"
_TOTALS = "E2E8F0"
_GRID = "CBD5E1"
_STATUS_COLORS = {
    "present": "15803D",
    "late": "B45309",
    "early_leave": "B45309",
    "absent": "B91C1C",
    "on_leave": "1D4ED8",
    "half_day": "6D28D9",
    # Security-alert severities (used by security_export's colored column).
    "critical": "B91C1C",
    "high": "C2410C",
    "medium": "B45309",
    "low": "64748B",
}

@dataclass(frozen=True)
class _ReportTable:
    """Single source of truth for a report's table shape, shared by the Excel
    and PDF renderers so columns can't drift between formats.

    Widths stay per-format only because the units differ (Excel character
    widths vs PDF relative column weights) — keeping both in the same spec
    makes a mismatch visible at a glance.
    """

    title: str
    sheet_name: str
    headers: tuple[str, ...]
    aligns: str  # per-column: L(eft) / C(enter) / R(ight)
    xlsx_widths: tuple[float, ...]
    pdf_widths: tuple[float, ...]
    status_col: int | None  # 0-based index of the status column (for colors)
    landscape: bool = False


_ATTENDANCE_TABLE = _ReportTable(
    title="Attendance Report",
    sheet_name="Attendance",
    headers=("Date", "Code", "Employee", "Status", "Check in", "Check out", "Worked", "Overtime"),
    aligns="LLLCCCRR",
    xlsx_widths=(13, 12, 28, 13, 11, 11, 11, 11),
    pdf_widths=(13, 10.5, 23, 13, 11.5, 12.5, 10, 11.5),
    status_col=3,
)


@dataclass(frozen=True)
class ExportContext:
    """Presentation metadata shared by every format."""

    org_name: str | None
    generated_at: datetime  # aware, UTC
    tz_offset: int  # JS getTimezoneOffset(): minutes, UTC minus local


def _viewer_tz(ctx: ExportContext) -> timezone:
    return timezone(timedelta(minutes=-ctx.tz_offset))


def _to_local(dt: datetime | None, ctx: ExportContext) -> datetime | None:
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(_viewer_tz(ctx))


def _fmt_time(dt: datetime | None, ctx: ExportContext) -> str:
    local = _to_local(dt, ctx)
    return local.strftime("%H:%M") if local else "-"


def _iso_local(dt: datetime | None, ctx: ExportContext) -> str:
    """ISO-8601 with the viewer's offset — machine-readable AND localized."""
    local = _to_local(dt, ctx)
    return local.isoformat() if local else ""


def _fmt_minutes(minutes: int) -> str:
    if not minutes:
        return "-"
    hours, rest = divmod(minutes, 60)
    if hours and rest:
        return f"{hours}h {rest:02d}m"
    return f"{hours}h" if hours else f"{rest}m"


def _status_label(status: str) -> str:
    return status.replace("_", " ").capitalize()


def _generated_line(ctx: ExportContext) -> str:
    local = ctx.generated_at.astimezone(_viewer_tz(ctx))
    sign = "-" if ctx.tz_offset > 0 else "+"
    off_h, off_m = divmod(abs(ctx.tz_offset), 60)
    stamp = local.strftime("%Y-%m-%d %H:%M")
    return f"Generated {stamp} (UTC{sign}{off_h:02d}:{off_m:02d}) - Attendance Platform"


def _attendance_subtitle(report: AttendanceRangeReport) -> str:
    start = report.period_start.strftime("%b %d, %Y")
    end = report.period_end.strftime("%b %d, %Y")
    return start if start == end else f"{start} - {end}"


def _attendance_kpis(report: AttendanceRangeReport) -> list[tuple[str, str]]:
    return [
        ("Records", str(report.total_records)),
        ("Present", str(report.present)),
        ("Late / early", str(report.late_or_early)),
        ("Hours worked", _fmt_minutes(report.worked_minutes)),
        ("Overtime", _fmt_minutes(report.overtime_minutes)),
    ]


def _attendance_rows(report: AttendanceRangeReport, ctx: ExportContext) -> list[list[str]]:
    return [
        [
            e.work_date.strftime("%b %d, %Y"),
            e.employee_code,
            e.employee_name,
            # The page badges attendance_type when set (e.g. a manual entry),
            # falling back to status — mirror that for the displayed label.
            # Colors are keyed on e.status separately (see status_keys).
            _status_label(e.attendance_type or e.status),
            _fmt_time(e.check_in_at, ctx),
            _fmt_time(e.check_out_at, ctx),
            _fmt_minutes(e.worked_minutes),
            _fmt_minutes(e.overtime_minutes),
        ]
        for e in report.entries
    ]


def _attendance_totals(report: AttendanceRangeReport) -> list[str]:
    n = report.total_records
    label = f"Totals - {n} record" + ("s" if n != 1 else "")
    return [
        "", "", label, "", "", "",
        _fmt_minutes(report.worked_minutes),
        _fmt_minutes(report.overtime_minutes),
    ]


# ── CSV ──────────────────────────────────────────────────────────────────────


def _csv_bytes(headers: list[str], rows: list[list[object]]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    # BOM so Excel auto-detects UTF-8 (names may be non-ASCII).
    return codecs.BOM_UTF8 + out.getvalue().encode("utf-8")


def _attendance_csv(report: AttendanceRangeReport, ctx: ExportContext) -> bytes:
    headers = [
        "Date", "Employee Code", "Employee Name", "Department", "Status", "Type",
        "Check In", "Check Out", "Worked Minutes", "Overtime Minutes",
        "Method In", "Method Out",
    ]
    rows = [
        [
            e.work_date.isoformat(),
            e.employee_code,
            e.employee_name,
            e.department or "",
            e.status,
            e.attendance_type or "",
            _iso_local(e.check_in_at, ctx),
            _iso_local(e.check_out_at, ctx),
            e.worked_minutes,
            e.overtime_minutes,
            e.check_in_method or "",
            e.check_out_method or "",
        ]
        for e in report.entries
    ]
    return _csv_bytes(headers, rows)


# ── Excel ────────────────────────────────────────────────────────────────────


@lru_cache(maxsize=1)
def _xl() -> SimpleNamespace:
    """openpyxl style objects, built on first export (lazy import)."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color=_GRID)
    accent = Side(style="medium", color=_BLUE)
    return SimpleNamespace(
        title=Font(name="Calibri", size=16, bold=True, color="FFFFFF"),
        subtitle=Font(name="Calibri", size=11, color="E2E8F0"),
        meta=Font(name="Calibri", size=9, italic=True, color=_SLATE),
        section=Font(name="Calibri", size=9, bold=True, color=_SLATE),
        kpi_label=Font(name="Calibri", size=9, bold=True, color=_SLATE),
        kpi_value=Font(name="Calibri", size=14, bold=True, color=_NAVY),
        th=Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
        total=Font(name="Calibri", size=11, bold=True, color=_NAVY),
        status_font=lambda color: Font(name="Calibri", size=11, bold=True, color=color),
        navy_fill=PatternFill("solid", fgColor=_NAVY),
        blue_fill=PatternFill("solid", fgColor=_BLUE),
        card_fill=PatternFill("solid", fgColor=_CARD),
        zebra_fill=PatternFill("solid", fgColor=_ZEBRA),
        totals_fill=PatternFill("solid", fgColor=_TOTALS),
        row_border=Border(bottom=thin),
        card_label_border=Border(left=accent, top=thin, right=thin),
        card_value_border=Border(left=accent, bottom=thin, right=thin),
        header_border=Border(bottom=accent),
        totals_border=Border(top=Side(style="medium", color=_NAVY)),
        align={
            "L": Alignment(horizontal="left", vertical="center"),
            "C": Alignment(horizontal="center", vertical="center"),
            "R": Alignment(horizontal="right", vertical="center"),
        },
    )


def _xlsx_scaffold(
    ws: "Worksheet",
    ncols: int,
    title: str,
    subtitle: str,
    ctx: ExportContext,
    kpis: list[tuple[str, str]],
) -> int:
    """Title band, accent strip, meta line and KPI cards. Returns the table header row."""
    from openpyxl.utils import get_column_letter

    S = _xl()
    last_col = get_column_letter(ncols)

    # Rows 1-2: navy title band; row 3: thin blue accent strip.
    ws.merge_cells(f"A1:{last_col}1")
    ws.merge_cells(f"A2:{last_col}2")
    band_subtitle = subtitle if not ctx.org_name else f"{subtitle}  |  {ctx.org_name}"
    ws["A1"], ws["A2"] = title, band_subtitle
    ws["A1"].font, ws["A2"].font = S.title, S.subtitle
    for row in (1, 2):
        for col in range(1, ncols + 1):
            ws.cell(row=row, column=col).fill = S.navy_fill
        ws.cell(row=row, column=1).alignment = S.align["L"]
    for col in range(1, ncols + 1):
        ws.cell(row=3, column=col).fill = S.blue_fill
    ws.row_dimensions[1].height = 27
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 3

    ws.merge_cells(f"A4:{last_col}4")
    ws["A4"] = _generated_line(ctx)
    ws["A4"].font = S.meta

    # KPI cards: label + value boxed with a blue accent on the left edge.
    ws.cell(row=6, column=1, value="SUMMARY").font = S.section
    for i, (label, value) in enumerate(kpis, start=1):
        label_cell = ws.cell(row=7, column=i, value=label.upper())
        value_cell = ws.cell(row=8, column=i, value=value)
        label_cell.font, value_cell.font = S.kpi_label, S.kpi_value
        label_cell.alignment = value_cell.alignment = S.align["L"]
        label_cell.fill = value_cell.fill = S.card_fill
        label_cell.border = S.card_label_border
        value_cell.border = S.card_value_border
    ws.row_dimensions[7].height = 14
    ws.row_dimensions[8].height = 20

    ws.cell(row=10, column=1, value="DETAIL").font = S.section
    return 11


def _xlsx_table(
    ws: "Worksheet",
    start_row: int,
    headers: list[str],
    rows: list[list[object]],
    aligns: str,
    widths: list[int],
    status_col: int | None,
    totals: list[object] | None,
    status_keys: list[str] | None = None,
) -> None:
    """Render the data table.

    ``status_keys`` carries the raw per-row status used for color lookup —
    independent of the displayed label, which may be a free-form
    ``attendance_type`` outside the status vocabulary.
    """
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    S = _xl()
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font, cell.fill = S.th, S.navy_fill
        cell.border = S.header_border
        cell.alignment = S.align[aligns[col - 1]]
    ws.row_dimensions[start_row].height = 20

    for r, row in enumerate(rows):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=start_row + 1 + r, column=col, value=value)
            cell.border = S.row_border
            cell.alignment = S.align[aligns[col - 1]]
            if r % 2 == 1:
                cell.fill = S.zebra_fill
            if status_col is not None and col == status_col:
                key = status_keys[r] if status_keys else str(value).lower().replace(" ", "_")
                cell.font = S.status_font(_STATUS_COLORS.get(key, _NAVY))

    last_data_row = start_row + len(rows)
    if totals is not None and rows:
        total_row = last_data_row + 1
        for col, value in enumerate(totals, start=1):
            cell = ws.cell(row=total_row, column=col, value=value)
            cell.font = S.total
            cell.fill = S.totals_fill
            cell.border = S.totals_border
            cell.alignment = S.align[aligns[col - 1]]

    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    # Filter/freeze cover the data only, never the totals row.
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{max(last_data_row, start_row)}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)

    # Print-ready: fit to page width, repeat the header row, footer with paging.
    ws.print_title_rows = f"{start_row}:{start_row}"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.oddFooter.left.text = "Attendance Platform"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8


def _xlsx_bytes(wb: "Workbook") -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _render_xlsx(
    spec: _ReportTable,
    subtitle: str,
    ctx: ExportContext,
    kpis: list[tuple[str, str]],
    rows: list[list[object]],
    totals: list[object] | None,
    status_keys: list[str] | None = None,
    percent_col: int | None = None,
) -> bytes:
    """Render one report worksheet from its shared table spec."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = spec.sheet_name
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if spec.landscape else "portrait"

    start = _xlsx_scaffold(ws, len(spec.headers), spec.title, subtitle, ctx, kpis)
    _xlsx_table(
        ws,
        start,
        list(spec.headers),
        rows,
        spec.aligns,
        list(spec.xlsx_widths),
        status_col=None if spec.status_col is None else spec.status_col + 1,
        totals=totals,
        status_keys=status_keys,
    )
    if percent_col is not None:
        total_rows = len(rows) + (1 if totals is not None and rows else 0)
        for r in range(total_rows):
            ws.cell(row=start + 1 + r, column=percent_col + 1).number_format = "0.0%"
    return _xlsx_bytes(wb)


# ── PDF ──────────────────────────────────────────────────────────────────────


@lru_cache(maxsize=1)
def _pdf() -> SimpleNamespace:
    """reportlab color objects, built on first export (lazy import)."""
    from reportlab.lib import colors

    return SimpleNamespace(
        navy=colors.HexColor(f"#{_NAVY}"),
        blue=colors.HexColor(f"#{_BLUE}"),
        card=colors.HexColor(f"#{_CARD}"),
        zebra=colors.HexColor(f"#{_ZEBRA}"),
        totals=colors.HexColor(f"#{_TOTALS}"),
        grid=colors.HexColor(f"#{_GRID}"),
        slate=colors.HexColor(f"#{_SLATE}"),
        band_text=colors.HexColor("#CBD5E1"),
        white=colors.white,
        hex=colors.HexColor,
    )


def _pdf_styles(aligns: str):
    """Per-column header / body / totals paragraph styles."""
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle

    C = _pdf()
    ta_map = {"L": TA_LEFT, "C": TA_CENTER, "R": TA_RIGHT}
    heads, bodies, totals = [], [], []
    for i, code in enumerate(aligns):
        ta = ta_map[code]
        heads.append(
            ParagraphStyle(
                f"h{i}", fontName="Helvetica-Bold", fontSize=8.5, leading=11,
                textColor=C.white, alignment=ta,
            )
        )
        bodies.append(
            ParagraphStyle(f"b{i}", fontName="Helvetica", fontSize=8.5, leading=11, alignment=ta)
        )
        totals.append(
            ParagraphStyle(
                f"t{i}", fontName="Helvetica-Bold", fontSize=8.5, leading=11,
                textColor=C.navy, alignment=ta,
            )
        )
    return heads, bodies, totals


def _pdf_kpi_cards(kpis: list[tuple[str, str]], usable: float):
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Table, TableStyle

    C = _pdf()
    label_style = ParagraphStyle(
        "kpiLabel", fontName="Helvetica-Bold", fontSize=7, textColor=C.slate
    )
    value_style = ParagraphStyle(
        "kpiValue", fontName="Helvetica-Bold", fontSize=14, textColor=C.navy
    )
    gap = 3 * mm
    card_w = (usable - gap * (len(kpis) - 1)) / len(kpis)

    cards = []
    for label, value in kpis:
        card = Table(
            [[Paragraph(label.upper(), label_style)], [Paragraph(value, value_style)]],
            colWidths=[card_w],
        )
        card.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), C.card),
                    ("BOX", (0, 0), (-1, -1), 0.6, C.grid),
                    ("LINEBEFORE", (0, 0), (0, -1), 2.2, C.blue),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, 0), 6),
                    ("BOTTOMPADDING", (0, 1), (-1, 1), 6),
                    ("TOPPADDING", (0, 1), (-1, 1), 1),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 1),
                ]
            )
        )
        cards.append(card)

    row: list[object] = []
    col_widths: list[float] = []
    for i, card in enumerate(cards):
        if i:
            row.append("")
            col_widths.append(gap)
        row.append(card)
        col_widths.append(card_w)
    holder = Table([row], colWidths=col_widths)
    holder.setStyle(
        TableStyle(
            [
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return holder


def _pdf_bytes(
    title: str,
    subtitle: str,
    ctx: ExportContext,
    kpis: list[tuple[str, str]],
    headers: list[str],
    rows: list[list[str]],
    col_widths: list[float],
    aligns: str,
    status_col: int | None,
    totals: list[str] | None,
    pagesize: tuple[float, float],
    status_keys: list[str] | None = None,
) -> bytes:
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdf_canvas
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    C = _pdf()
    buf = io.BytesIO()
    page_w = pagesize[0]
    margin = 15 * mm
    usable = page_w - 2 * margin
    footer_text = _generated_line(ctx)

    def draw_band(c: pdf_canvas.Canvas, first_page: bool) -> None:
        c.saveState()
        if first_page:
            band_h = 24 * mm
            c.setFillColor(C.navy)
            c.rect(0, pagesize[1] - band_h, page_w, band_h, stroke=0, fill=1)
            c.setFillColor(C.blue)
            c.rect(0, pagesize[1] - band_h - 1.2 * mm, page_w, 1.2 * mm, stroke=0, fill=1)
            c.setFillColor(C.white)
            c.setFont("Helvetica-Bold", 16)
            c.drawString(margin, pagesize[1] - 12 * mm, title)
            c.setFillColor(C.band_text)
            c.setFont("Helvetica", 10)
            c.drawString(margin, pagesize[1] - 18 * mm, subtitle)
            if ctx.org_name:
                c.setFont("Helvetica-Bold", 10)
                c.drawRightString(page_w - margin, pagesize[1] - 12 * mm, ctx.org_name)
        else:
            c.setFillColor(C.slate)
            c.setFont("Helvetica", 8)
            c.drawString(margin, pagesize[1] - 10 * mm, f"{title} - {subtitle}")
            c.setStrokeColor(C.grid)
            c.line(margin, pagesize[1] - 12 * mm, page_w - margin, pagesize[1] - 12 * mm)
        c.restoreState()

    class NumberedCanvas(pdf_canvas.Canvas):
        """Two-pass canvas so the footer can show 'Page X of Y'."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved: list[dict] = []

        def showPage(self):  # noqa: N802 (reportlab API)
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved)
            for i, state in enumerate(self._saved, start=1):
                self.__dict__.update(state)
                draw_band(self, first_page=(i == 1))
                self.setStrokeColor(C.grid)
                self.line(margin, 13 * mm, page_w - margin, 13 * mm)
                self.setFillColor(C.slate)
                self.setFont("Helvetica", 8)
                self.drawString(margin, 9 * mm, footer_text)
                self.drawRightString(page_w - margin, 9 * mm, f"Page {i} of {total}")
                super().showPage()
            super().save()

    doc = SimpleDocTemplate(
        buf,
        pagesize=pagesize,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=32 * mm,
        bottomMargin=20 * mm,
        title=f"{title} - {subtitle}",
        author="Attendance Platform",
    )

    section = ParagraphStyle(
        "section", fontName="Helvetica-Bold", fontSize=7.5, textColor=C.slate, spaceAfter=4
    )

    heads, bodies, total_styles = _pdf_styles(aligns)
    table_data = [[Paragraph(h, heads[i]) for i, h in enumerate(headers)]]
    table_data += [[Paragraph(str(v), bodies[i]) for i, v in enumerate(row)] for row in rows]
    has_totals = totals is not None and bool(rows)
    if has_totals:
        table_data.append([Paragraph(str(v), total_styles[i]) for i, v in enumerate(totals)])

    last_body = -2 if has_totals else -1
    style: list[tuple] = [
        ("BACKGROUND", (0, 0), (-1, 0), C.navy),
        ("LINEBELOW", (0, 0), (-1, 0), 1.2, C.blue),
        ("ROWBACKGROUNDS", (0, 1), (-1, last_body), [C.white, C.zebra]),
        ("LINEBELOW", (0, 1), (-1, last_body), 0.4, C.grid),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]
    if has_totals:
        style += [
            ("BACKGROUND", (0, -1), (-1, -1), C.totals),
            ("LINEABOVE", (0, -1), (-1, -1), 1, C.navy),
        ]
    if status_col is not None:
        for r in range(1, len(rows) + 1):
            # Color by the raw status key, not the (possibly free-form) label.
            key = (
                status_keys[r - 1]
                if status_keys
                else rows[r - 1][status_col].lower().replace(" ", "_")
            )
            color = _STATUS_COLORS.get(key)
            if color:
                style.append(("TEXTCOLOR", (status_col, r), (status_col, r), C.hex(f"#{color}")))

    total_w = sum(col_widths)
    data_table = Table(
        table_data, colWidths=[w / total_w * usable for w in col_widths], repeatRows=1
    )
    data_table.setStyle(TableStyle(style))

    empty_note = Paragraph(
        "No records for this period.",
        ParagraphStyle("empty", fontName="Helvetica-Oblique", fontSize=9, textColor=C.slate),
    )
    story = [
        Paragraph("SUMMARY", section),
        _pdf_kpi_cards(kpis, usable),
        Spacer(0, 7 * mm),
        Paragraph("DETAIL", section),
        data_table if rows else empty_note,
    ]
    doc.build(story, canvasmaker=NumberedCanvas)
    return buf.getvalue()


def _render_pdf(
    spec: _ReportTable,
    subtitle: str,
    ctx: ExportContext,
    kpis: list[tuple[str, str]],
    rows: list[list[str]],
    totals: list[str] | None,
    status_keys: list[str] | None = None,
) -> bytes:
    """Render one report PDF from its shared table spec."""
    from reportlab.lib.pagesizes import A4, landscape

    return _pdf_bytes(
        spec.title,
        subtitle,
        ctx,
        kpis,
        list(spec.headers),
        rows,
        col_widths=list(spec.pdf_widths),
        aligns=spec.aligns,
        status_col=spec.status_col,
        totals=totals,
        pagesize=landscape(A4) if spec.landscape else A4,
        status_keys=status_keys,
    )


# ── Report assemblies ────────────────────────────────────────────────────────
# The report wires its spec + row builders into both renderers side by side.


def _attendance_xlsx(report: AttendanceRangeReport, ctx: ExportContext) -> bytes:
    return _render_xlsx(
        _ATTENDANCE_TABLE, _attendance_subtitle(report), ctx, _attendance_kpis(report),
        _attendance_rows(report, ctx), _attendance_totals(report),
        status_keys=[e.status for e in report.entries],
    )


def _attendance_pdf(report: AttendanceRangeReport, ctx: ExportContext) -> bytes:
    return _render_pdf(
        _ATTENDANCE_TABLE, _attendance_subtitle(report), ctx, _attendance_kpis(report),
        _attendance_rows(report, ctx), _attendance_totals(report),
        status_keys=[e.status for e in report.entries],
    )


# ── Public API ───────────────────────────────────────────────────────────────


def export_attendance(report: AttendanceRangeReport, fmt: str, ctx: ExportContext) -> bytes:
    renderers = {"csv": _attendance_csv, "xlsx": _attendance_xlsx, "pdf": _attendance_pdf}
    return renderers[fmt](report, ctx)
